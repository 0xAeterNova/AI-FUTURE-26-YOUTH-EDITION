"""Small-corpus ingestion with explicit physical-page references."""
from __future__ import annotations
import hashlib
import json
import re
from pathlib import Path
from typing import Any

VALID = {".txt", ".md", ".pdf", ".docx"}


def extract(path: Path) -> list[str]:
    ext = path.suffix.lower()
    if ext in {".txt", ".md"}:
        return [path.read_text(encoding="utf-8-sig")]
    if ext == ".pdf":
        try:
            import pymupdf
        except ImportError:
            try:
                import fitz as pymupdf
            except ImportError as e:
                raise RuntimeError("Install PDF support: pip install '.[documents]'") from e
        pdf = pymupdf.open(str(path))
        try:
            return [p.get_text("text", sort=True) for p in pdf]
        finally:
            pdf.close()
    if ext == ".docx":
        try:
            from docx import Document
        except ImportError as e:
            raise RuntimeError("Install DOCX support: pip install '.[documents]'") from e
        doc = Document(path)
        out = [p.text for p in doc.paragraphs if p.text.strip()]
        for tab in doc.tables:
            for row in tab.rows:
                out.append(" | ".join(c.text.replace("\n", " ؛ ") for c in row.cells))
        return ["\n".join(out)]  # DOCX lacks a reliable physical-page mapping
    raise ValueError(f"Unsupported extension: {ext}")


def title_from(path: Path, page: str) -> str:
    lines = [line.strip() for line in page.splitlines() if line.strip()]
    # Filenames frequently carry the service name; first heading is a fallback for numbered filenames.
    title = path.stem.replace("_", " ").replace("-", " ").strip()
    if (not title or title.isdigit() or bool(re.match(r"^\d+\s", title)) or title.lower().startswith(("doc", "file", "service"))) and lines:
        title = next((s for s in lines[:7] if 5 < len(s) < 100), lines[0][:100])
    return title.lstrip("# ").strip()


def split_chunks(text: str, target: int = 660) -> list[str]:
    """Preserve nearby lines and bullets; split enormous paragraphs at sentence boundaries."""
    import re
    lines = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if len(line) > target:
            segments = re.split(r"(?<=[.!؟؛:])\s+", line)
            if len(segments) == 1:
                segments = [line[i:i+target] for i in range(0, len(line), target)]
            lines.extend(segments)
        else:
            lines.append(line)
    if not lines:
        return []
    chunks, current = [], []
    for ln in lines:
        if current and sum(map(len, current)) + len(ln) > target:
            chunks.append("\n".join(current))
            current = current[-1:] if len(current[-1]) < 150 else []  # small overlap for boundaries
        current.append(ln)
    if current:
        chunks.append("\n".join(current))
    return chunks


def index_folder(folder: str, output: str, manifest: str | None = None, *, ocr: str = "auto", use_markitdown: bool = True) -> dict[str, Any]:
    root = Path(folder)
    if not root.is_dir():
        raise FileNotFoundError(f"Document directory not found: {folder}")
    paths = sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in VALID)
    if not paths:
        raise RuntimeError(f"No .pdf/.docx/.txt/.md files in {folder}")
    result: dict[str, Any] = {"version": 1, "documents": [], "chunks": [], "warnings": [], "extraction_reports": []}
    cache = {}
    sources = []
    from .quality import extract_with_audit
    def get_pages(path):
        if path not in cache:
            pages, reports, warnings = extract_with_audit(path, ocr=ocr, use_markitdown=use_markitdown)
            cache[path] = pages
            result["extraction_reports"].append({"file": str(path), "pages": reports})
            result["warnings"].extend(warnings)
        return cache[path]
    if manifest:
        config = json.loads(Path(manifest).read_text(encoding="utf-8"))
        for item in config.get("services", []):
            path = (root / item["file"]).resolve()
            if not path.is_file() or path.suffix.lower() not in VALID or not path.is_relative_to(root.resolve()):
                raise ValueError(f"Invalid manifest file path: {item['file']}")
            all_pages = get_pages(path)
            which = item.get("pages", list(range(1, len(all_pages) + 1)))
            if not which or any(not isinstance(i, int) or not 1 <= i <= len(all_pages) for i in which):
                raise ValueError(f"Invalid 1-based pages in manifest for {path.name}")
            sources.append((path, [(i, all_pages[i-1]) for i in which], item.get("title", "").strip()))
        if not sources:
            raise ValueError("Manifest has no services")
    else:
        for path in paths:
            sources.append((path, list(enumerate(get_pages(path), 1)), ""))
    for n, (path, numbered_pages, custom_title) in enumerate(sources, 1):
        did = f"D{n:02d}"
        title = custom_title or title_from(path, next((p for _, p in numbered_pages if p.strip()), ""))
        doc = {"id": did, "title": title, "file": str(path.resolve()), "pages": [], "digest": hashlib.sha256(path.read_bytes()).hexdigest()[:16]}
        for i, content in numbered_pages:
            page_text = content.strip()
            pid = f"{did}-P{i}"
            doc["pages"].append({"id": pid, "number": i, "text": page_text})
            for part, excerpt in enumerate(split_chunks(page_text), 1):
                result["chunks"].append({"id": f"{pid}-C{part}", "page_id": pid, "doc_id": did, "text": excerpt})
            if len(page_text) < 80:
                result["warnings"].append(f"{path.name} page {i}: only {len(page_text)} extracted characters; check scanned/blank page")
        result["documents"].append(doc)
    from .ledger import build_ledger
    result["ledger"] = build_ledger(result)
    dest = Path(output)
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def load_index(path: str) -> dict[str, Any]:
    obj = json.loads(Path(path).read_text(encoding="utf-8"))
    if obj.get("version") != 1:
        raise ValueError("Unsupported index version")
    return obj
